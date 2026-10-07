#!/usr/bin/env python3
"""Generates the workbooks of this directory with openpyxl (and post-processes a few to
carry what openpyxl cannot write: an extLst, a data descriptor). Usage:

    python3 generate.py OUT_DIR

LibreOffice versions of them are made with
    soffice --headless --convert-to xlsx --outdir OUT_DIR/lo OUT_DIR/*.xlsx
which also gives every formula a cached value.
"""
import datetime as dt
import sys
import zipfile
from pathlib import Path

import openpyxl
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import CellIsRule, DataBarRule
from openpyxl.styles import Font, PatternFill
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.utils.datetime import CALENDAR_MAC_1904
from openpyxl.worksheet.table import Table, TableStyleInfo

out = Path(sys.argv[1])
out.mkdir(parents=True, exist_ok=True)


def rows(ws, header, data):
    ws.append(header)
    for r in data:
        ws.append(r)


# 1. Several sheets, strings, dates, formulas, merged cells, a hidden sheet, an empty
#    sheet, a title row sheet, a chart sheet, frozen panes, an autofilter, a table.
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Sales"
rows(ws, ["id", "region", "amount", "when", "paid", "note"], [
    [1, "North", 10.5, dt.datetime(2021, 1, 1), True, "first"],
    [2, "South", 20.25, dt.datetime(2021, 1, 2, 13, 30), False, None],
    [3, "North & <East>", 7, dt.date(2021, 3, 1), True, "tab\there"],
    [4, "Ünïcode ✓", 1e-7, dt.datetime(2020, 2, 29), None, "quote \" '"],
])
ws["G1"] = "total"
ws["G2"] = "=SUM(C2:C5)"
ws.freeze_panes = "A2"
ws.merge_cells("A8:C9")
ws["A8"] = "merged text"
tab = Table(displayName="SalesTable", ref="A1:F5")
tab.tableStyleInfo = TableStyleInfo(name="TableStyleMedium9", showRowStripes=True)
ws.add_table(tab)
ws2 = wb.create_sheet("Lookup")
rows(ws2, ["code", "label"], [["A", "alpha"], ["B", "beta"]])
ws2.sheet_state = "hidden"
wb.create_sheet("Empty")
ws3 = wb.create_sheet("Report")
ws3["A1"] = "Quarterly report"
ws3.append(["id", "name"])
ws3.append([1, "a"])
ws4 = wb.create_sheet("Data")
rows(ws4, ["k", "v"], [["x", 1], ["y", 2], ["z", 3]])
ws4.auto_filter.ref = "A1:B4"
cs = wb.create_chartsheet("Chart")
chart = BarChart()
chart.add_data(Reference(ws4, min_col=2, min_row=1, max_row=4), titles_from_data=True)
cs.add_chart(chart)
wb.save(out / "multi.xlsx")

# 2. Rich text and inline-looking strings.
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Rich"
ws.append(["plain", "rich", "empty string", "spaces"])
ws.append(["a", CellRichText("bold ", TextBlock(InlineFont(b=True), "BOLD"), " tail"), "", "  padded  "])
wb.save(out / "rich.xlsx")

# 3. Dates in the 1904 system.
wb = openpyxl.Workbook()
wb.epoch = CALENDAR_MAC_1904
ws = wb.active
ws.title = "Mac"
rows(ws, ["day", "stamp"], [[dt.datetime(2021, 1, 1), dt.datetime(2021, 1, 1, 6, 0)], [dt.datetime(1999, 12, 31), dt.datetime(2000, 1, 1, 0, 0, 1)]])
wb.save(out / "dates1904.xlsx")

# 4. A wide sheet, styles with many formats, conditional formatting.
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Wide"
ws.append([f"col{i}" for i in range(120)])
for r in range(6):
    ws.append([r * 1000 + c for c in range(120)])
st = wb.create_sheet("Styled")
st.append(["n", "x"])
for i in range(60):
    st.append([i, i * 1.5])
    st.cell(row=i + 2, column=2).number_format = f'0.{"0" * (i % 8)}" u{i}"'
    st.cell(row=i + 2, column=1).font = Font(bold=i % 2 == 0, color="FF0000")
    st.cell(row=i + 2, column=1).fill = PatternFill("solid", fgColor=f"{i % 256:02X}AA55")
st.conditional_formatting.add("A2:A60", CellIsRule(operator="greaterThan", formula=["100"], fill=PatternFill("solid", bgColor="FFFF00")))
st.conditional_formatting.add("B2:B60", DataBarRule(start_type="min", end_type="max", color="638EC6"))
wb.save(out / "wide_styled.xlsx")

# 5. A 17-digit id column and a Persian header.
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "ids"
ws.append(["id", "می‌خواهم", "mixed"])
ws.append([12345678901234567, "الف", 1])
ws.append([12345678901234568, "ب", 2.5])
ws.append([-9007199254740993, "پ", "text"])
wb.save(out / "ids_persian.xlsx")


def inject(path, edits, descriptor=False):
    """Rewrites parts of an xlsx: edits maps a part name to a function of its bytes."""
    src = zipfile.ZipFile(path)
    tmp = path.with_suffix(".tmp")
    with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as dst:
        for info in src.infolist():
            data = src.read(info.filename)
            if info.filename in edits:
                data = edits[info.filename](data)
            dst.writestr(info.filename, data)
    tmp.replace(path)


# 6. An extLst as Excel writes it (x14 conditional formatting and sparklines): elements
#    in other namespaces, one of them named like a worksheet element.
EXT = (
    b'<extLst><ext uri="{78C0D931-6437-407d-A8EE-F0AAD7539E65}" '
    b'xmlns:x14="http://schemas.microsoft.com/office/spreadsheetml/2009/9/main">'
    b'<x14:conditionalFormattings><x14:conditionalFormatting '
    b'xmlns:xm="http://schemas.microsoft.com/office/excel/2006/main">'
    b'<x14:cfRule type="dataBar" id="{A}"><x14:dataBar minLength="0" maxLength="100">'
    b'<x14:cfvo type="autoMin"/><x14:cfvo type="autoMax"/></x14:dataBar></x14:cfRule>'
    b'<xm:sqref>B2:B5</xm:sqref></x14:conditionalFormatting></x14:conditionalFormattings></ext>'
    b'<ext uri="{05C60535-1F16-4fd2-B633-F4F36F0B64E0}" '
    b'xmlns:x14="http://schemas.microsoft.com/office/spreadsheetml/2009/9/main" '
    b'xmlns:xm="http://schemas.microsoft.com/office/excel/2006/main">'
    b'<x14:sparklineGroups><x14:sparklineGroup><x14:sparklines><x14:sparkline>'
    b'<xm:f>Sheet1!A2:A5</xm:f><xm:sqref>C2</xm:sqref></x14:sparkline></x14:sparklines>'
    b'</x14:sparklineGroup></x14:sparklineGroups></ext></extLst>'
)
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Ext"
rows(ws, ["a", "b", "c"], [[1, 2, 3], [4, 5, 6], [7, 8, 9], [10, 11, 12]])
wb.save(out / "extlst.xlsx")
inject(out / "extlst.xlsx", {"xl/worksheets/sheet1.xml": lambda d: d.replace(b"</worksheet>", EXT + b"</worksheet>")})

# 7. Namespace-prefixed worksheet elements, as some writers produce.
def prefixed(data):
    import re
    s = data.decode()
    s = re.sub(r"<(/?)(worksheet|sheetData|row|c|v|is|t)([ >/])", r"<\1x:\2\3", s)
    s = s.replace('<x:worksheet ', '<x:worksheet xmlns:x="http://schemas.openxmlformats.org/spreadsheetml/2006/main" ', 1)
    s = s.replace(' xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"', "", 1)
    return s.encode()

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Prefixed"
rows(ws, ["a", "b"], [[1, "x"], [2, "y"]])
wb.save(out / "prefixed.xlsx")
inject(out / "prefixed.xlsx", {"xl/worksheets/sheet1.xml": prefixed})


# 8. The same workbook through writers that stream: a data descriptor after each part
#    (an unseekable output), and zip64 extra fields in the local headers.
class Unseekable:
    def __init__(self, fh):
        self.fh = fh

    def write(self, b):
        return self.fh.write(b)

    def flush(self):
        self.fh.flush()

    def seekable(self):
        return False

    def tell(self):
        return self.fh.tell()


def restream(src_path, dst_path, descriptor=False, zip64=False):
    src = zipfile.ZipFile(src_path)
    with open(dst_path, "wb") as fh:
        target = Unseekable(fh) if descriptor else fh
        with zipfile.ZipFile(target, "w", zipfile.ZIP_DEFLATED, allowZip64=True) as dst:
            for info in src.infolist():
                with dst.open(info.filename, "w", force_zip64=zip64) as part:
                    part.write(src.read(info.filename))


restream(out / "multi.xlsx", out / "descriptor.xlsx", descriptor=True)
restream(out / "multi.xlsx", out / "zip64_local.xlsx", zip64=True)
